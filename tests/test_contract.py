import os, shutil
import pytest
import docx
import laptops

HAS_SOFFICE = bool(shutil.which("soffice") or shutil.which("libreoffice"))


@pytest.fixture
def assignment(conn):
    s = conn.execute("INSERT INTO students(name,student_no,id_number) VALUES('Thandi Nkosi','2024001','9001015800087')").lastrowid
    l = conn.execute("INSERT INTO laptops(asset_tag,model,serial) VALUES('LT-001','Dell Latitude 3420','SN12345')").lastrowid
    a = conn.execute("INSERT INTO assignments(student_id,laptop_id,issued_on) VALUES(?,?,'2026-03-15')", (s, l)).lastrowid
    conn.commit()
    return a


@pytest.fixture
def no_pdf(monkeypatch):
    monkeypatch.setattr(laptops, "to_pdf", lambda *a, **k: False)


def docx_path():
    return os.path.join(laptops.OUT, "loan_agreement_2024001_LT-001.docx")


def test_contract_docx_fields(assignment, no_pdf, capsys):
    laptops.make_contract(assignment)
    assert "Word file only" in capsys.readouterr().out
    assert os.path.exists(docx_path())
    d = docx.Document(docx_path())
    who, sign = d.tables
    assert who.rows[0].cells[1].text == "Thandi Nkosi"
    assert who.rows[1].cells[1].text == "2024001"
    assert who.rows[2].cells[1].text == "9001015800087"
    assert sign.rows[1].cells[1].text == "Thandi Nkosi"
    assert sign.rows[3].cells[1].text == "2026-03-15"
    texts = [p.text for p in d.paragraphs]
    assert "That the student received a Dell Latitude 3420" in texts
    assert "Serial Number: SN12345" in texts


def test_contract_signature_cell_blank(assignment, no_pdf):
    laptops.make_contract(assignment)
    sign = docx.Document(docx_path()).tables[1]
    assert sign.rows[2].cells[0].text.strip().startswith("Signature")
    assert sign.rows[2].cells[1].text == ""


def test_contract_old_template_values_gone(assignment, no_pdf):
    laptops.make_contract(assignment)
    d = docx.Document(docx_path())
    full = "\n".join(p.text for p in d.paragraphs)
    assert "Lenovo V15" not in full and "PF5X3HNC" not in full


def test_contract_missing_id_number_is_blank(conn, no_pdf):
    s = conn.execute("INSERT INTO students(name,student_no) VALUES('No Id','77')").lastrowid
    l = conn.execute("INSERT INTO laptops(asset_tag,model,serial) VALUES('LT-9','M','S')").lastrowid
    a = conn.execute("INSERT INTO assignments(student_id,laptop_id,issued_on) VALUES(?,?,'2026-01-01')", (s, l)).lastrowid
    conn.commit()
    laptops.make_contract(a)
    who = docx.Document(os.path.join(laptops.OUT, "loan_agreement_77_LT-9.docx")).tables[0]
    assert who.rows[2].cells[1].text == ""


def test_contract_unknown_assignment(no_pdf):
    with pytest.raises(SystemExit):
        laptops.make_contract(999)


def test_contract_written_only_to_tmp(assignment, no_pdf, tmp_env):
    laptops.make_contract(assignment)
    assert docx_path().startswith(str(tmp_env))


def test_contract_regenerate_overwrites(assignment, no_pdf):
    laptops.make_contract(assignment)
    laptops.make_contract(assignment)
    assert len(os.listdir(laptops.OUT)) == 1


@pytest.mark.skipif(not HAS_SOFFICE, reason="LibreOffice (soffice) not available")
def test_contract_pdf_conversion(assignment, capsys):
    laptops.make_contract(assignment)
    pdf = docx_path()[:-5] + ".pdf"
    assert os.path.exists(pdf) and os.path.getsize(pdf) > 0
    with open(pdf, "rb") as f:
        assert f.read(4) == b"%PDF"
    assert "Contract:" in capsys.readouterr().out


def test_to_pdf_returns_false_without_converters(tmp_path, monkeypatch):
    monkeypatch.setattr(shutil, "which", lambda *_: None)
    monkeypatch.setattr(os.path, "exists", lambda p: False)
    import builtins
    real = builtins.__import__
    def fake(name, *a, **k):
        if name == "docx2pdf":
            raise ImportError
        return real(name, *a, **k)
    monkeypatch.setattr(builtins, "__import__", fake)
    assert laptops.to_pdf(str(tmp_path / "x.docx"), str(tmp_path / "x.pdf")) is False
