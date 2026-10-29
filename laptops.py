#!/usr/bin/env python3
"""Laptop loan tracker: assign laptops to students and generate contracts."""
import argparse, csv, io, os, re, sqlite3, subprocess, sys
from datetime import date

DB = os.path.join(os.path.dirname(os.path.abspath(__file__)), "laptops.db")
OUT = os.path.join(os.path.dirname(os.path.abspath(__file__)), "contracts")
TEMPLATE = os.path.join(os.path.dirname(os.path.abspath(__file__)), "Loan_device_Form.docx")

SCHEMA = """
CREATE TABLE IF NOT EXISTS students(
  id INTEGER PRIMARY KEY, name TEXT NOT NULL, student_no TEXT UNIQUE,
  grade TEXT, guardian_name TEXT, guardian_email TEXT, id_number TEXT);
CREATE TABLE IF NOT EXISTS laptops(
  id INTEGER PRIMARY KEY, asset_tag TEXT UNIQUE NOT NULL, model TEXT,
  serial TEXT, value REAL DEFAULT 0);
CREATE TABLE IF NOT EXISTS assignments(
  id INTEGER PRIMARY KEY,
  student_id INTEGER NOT NULL REFERENCES students(id),
  laptop_id INTEGER NOT NULL REFERENCES laptops(id),
  issued_on TEXT NOT NULL, returned_on TEXT, signed_on TEXT);
-- a laptop can only be out on one open assignment at a time
CREATE UNIQUE INDEX IF NOT EXISTS one_open_per_laptop
  ON assignments(laptop_id) WHERE returned_on IS NULL;
"""

def db():
    c = sqlite3.connect(DB); c.row_factory = sqlite3.Row
    c.execute("PRAGMA foreign_keys=ON"); c.executescript(SCHEMA)
    if "id_number" not in [r[1] for r in c.execute("PRAGMA table_info(students)")]:
        c.execute("ALTER TABLE students ADD COLUMN id_number TEXT")
    return c

def find(c, table, col, val):
    r = c.execute(f"SELECT * FROM {table} WHERE {col}=?", (val,)).fetchone()
    if not r: sys.exit(f"No {table[:-1]} with {col}={val}")
    return r

def cmd_add_student(a):
    with db() as c:
        c.execute("INSERT INTO students(name,student_no,grade,guardian_name,guardian_email,id_number) VALUES(?,?,?,?,?,?)",
                  (a.name, a.student_no, a.grade, a.guardian, a.guardian_email, a.id_number))
    print("Added", a.name)

ALIASES = {"name": {"name", "student name", "student name and surname", "full name", "name and surname", "name surname"},
           "student_no": {"student_no", "student number", "student no", "studentno", "student_number"},
           "id_number": {"id_number", "id", "id number", "id/passport", "student id/passport", "passport", "id_passport"}}

LAPTOP_ALIASES = {"tag": {"asset_tag", "asset tag", "tag", "asset", "asset no", "asset number"},
                  "model": {"model", "device", "device model"},
                  "serial": {"serial", "serial number", "serial no", "serial_number", "sn"}}

def read_csv(text, aliases, required):
    """Parse CSV text -> (list of dict rows with row numbers, error). Matches headers via aliases."""
    text = text.lstrip("\ufeff")
    try: dialect = csv.Sniffer().sniff(text[:2000], delimiters=",;\t")
    except csv.Error: dialect = csv.excel
    rows = list(csv.reader(io.StringIO(text), dialect))
    if not rows: return [], "The file is empty."
    head = [h.strip().lower() for h in rows[0]]
    col = {k: next((i for i, h in enumerate(head) if h in v), None) for k, v in aliases.items()}
    missing = [k for k in required if col[k] is None]
    if missing:
        return [], "Header row is missing: " + ", ".join(missing) + ". Found: " + ", ".join(rows[0])
    out = []
    for n, r in enumerate(rows[1:], start=2):
        if not any(x.strip() for x in r): continue
        out.append((n, {k: (r[i].strip() if i is not None and i < len(r) else "") for k, i in col.items()}))
    return out, None

def import_students(text):
    """Add students from CSV text. Returns (added, skipped, problems). Duplicates by student number are skipped."""
    rows, err = read_csv(text, ALIASES, ["name", "student_no"])
    if err: return 0, [], [err]
    added, skipped, problems = 0, [], []
    with db() as c:
        for n, r in rows:
            if not r["name"] or not r["student_no"]:
                problems.append(f"Row {n}: missing name or student number."); continue
            try:
                c.execute("INSERT INTO students(name,student_no,id_number) VALUES(?,?,?)",
                          (r["name"], r["student_no"], r["id_number"]))
                added += 1
            except sqlite3.IntegrityError:
                skipped.append(f"{r['name']} ({r['student_no']})")
    return added, skipped, problems

def import_laptops(text):
    """Add laptops from CSV text. Returns (added, skipped, problems). Duplicates by asset tag are skipped."""
    rows, err = read_csv(text, LAPTOP_ALIASES, ["tag", "serial"])
    if err: return 0, [], [err]
    added, skipped, problems = 0, [], []
    with db() as c:
        for n, r in rows:
            if not r["tag"] or not r["serial"]:
                problems.append(f"Row {n}: missing asset tag or serial number."); continue
            try:
                c.execute("INSERT INTO laptops(asset_tag,model,serial) VALUES(?,?,?)", (r["tag"], r["model"], r["serial"]))
                added += 1
            except sqlite3.IntegrityError:
                skipped.append(r["tag"])
    return added, skipped, problems

def cmd_import(a, fn=None, what="students"):
    added, skipped, problems = (fn or import_students)(open(a.file, encoding="utf-8-sig", newline="").read())
    print(f"Added {added} {what}.")
    if skipped: print(f"Skipped {len(skipped)} already in database:", ", ".join(skipped))
    for p in problems: print("Problem:", p)

def cmd_import_laptops(a): cmd_import(a, import_laptops, "laptops")

def cmd_add_laptop(a):
    with db() as c:
        c.execute("INSERT INTO laptops(asset_tag,model,serial,value) VALUES(?,?,?,?)",
                  (a.tag, a.model, a.serial, a.value))
    print("Added", a.tag)

def cmd_assign(a):
    with db() as c:
        s = find(c, "students", "student_no", a.student_no)
        l = find(c, "laptops", "asset_tag", a.tag)
        try:
            cur = c.execute("INSERT INTO assignments(student_id,laptop_id,issued_on) VALUES(?,?,?)",
                            (s["id"], l["id"], a.date or date.today().isoformat()))
        except sqlite3.IntegrityError:
            sys.exit(f"{a.tag} is already assigned. Return it first.")
    print(f"Assignment #{cur.lastrowid}: {l['asset_tag']} -> {s['name']}")
    make_contract(cur.lastrowid)

def cmd_return(a):
    with db() as c:
        n = c.execute("UPDATE assignments SET returned_on=? WHERE returned_on IS NULL AND laptop_id="
                      "(SELECT id FROM laptops WHERE asset_tag=?)",
                      (a.date or date.today().isoformat(), a.tag)).rowcount
    print("Returned" if n else "That laptop is not currently assigned.")

def cmd_signed(a):
    with db() as c:
        c.execute("UPDATE assignments SET signed_on=? WHERE id=?", (a.date or date.today().isoformat(), a.id))
    print("Marked signed.")

def cmd_list(a):
    q = """SELECT a.id, l.asset_tag, l.model, s.name, s.grade, a.issued_on,
                   COALESCE(a.returned_on,'-') ret, CASE WHEN a.signed_on IS NULL THEN 'NO' ELSE 'yes' END signed
           FROM assignments a JOIN students s ON s.id=a.student_id JOIN laptops l ON l.id=a.laptop_id
           {} ORDER BY a.id"""
    with db() as c:
        rows = c.execute(q.format("" if a.all else "WHERE a.returned_on IS NULL")).fetchall()
    print(f"{'ID':<4}{'TAG':<10}{'MODEL':<18}{'STUDENT':<20}{'GR':<4}{'ISSUED':<12}{'RETURNED':<12}SIGNED")
    for r in rows:
        print(f"{r['id']:<4}{r['asset_tag']:<10}{(r['model'] or ''):<18}{r['name']:<20}{(r['grade'] or ''):<4}"
              f"{r['issued_on']:<12}{r['ret']:<12}{r['signed']}")

def cmd_available(a):
    with db() as c:
        for r in c.execute("""SELECT asset_tag, model FROM laptops WHERE id NOT IN
                              (SELECT laptop_id FROM assignments WHERE returned_on IS NULL)"""):
            print(r["asset_tag"], r["model"] or "")

def cmd_contract(a): make_contract(a.id)

def set_runs(par, text):
    """Replace a paragraph's text, keeping the first run's formatting."""
    runs = par.runs
    runs[0].text = text
    for r in runs[1:]: r.text = ""

def fill_cell(cell, text):
    par = cell.paragraphs[0]
    if par.runs: par.runs[0].text = text
    else: par.add_run(text)

def make_contract(aid):
    import docx
    with db() as c:
        r = c.execute("""SELECT a.id, a.issued_on, s.name, s.student_no, s.id_number,
                                l.asset_tag, l.model, l.serial
                         FROM assignments a JOIN students s ON s.id=a.student_id
                         JOIN laptops l ON l.id=a.laptop_id WHERE a.id=?""", (aid,)).fetchone()
    if not r: sys.exit("No such assignment")
    d = docx.Document(TEMPLATE)
    who, sign = d.tables
    fill_cell(who.rows[0].cells[1], r["name"])
    fill_cell(who.rows[1].cells[1], r["student_no"] or "")
    fill_cell(who.rows[2].cells[1], r["id_number"] or "")
    fill_cell(sign.rows[1].cells[1], r["name"])          # Full name (print); signature left blank
    fill_cell(sign.rows[3].cells[1], r["issued_on"])
    for p in d.paragraphs:
        if p.text.startswith("That the student received a"):
            set_runs(p, f"That the student received a {r['model'] or ''}")
        elif p.text.startswith("Serial Number"):
            set_runs(p, f"Serial Number: {r['serial'] or ''}")
    os.makedirs(OUT, exist_ok=True)
    safe = lambda v: re.sub(r"[^A-Za-z0-9._-]", "_", str(v))   # no path separators in file names
    base = os.path.join(OUT, f"loan_agreement_{safe(r['student_no'])}_{safe(r['asset_tag'])}")
    d.save(base + ".docx")
    if to_pdf(base + ".docx", base + ".pdf"):
        print("Contract:", base + ".pdf")
    else:
        print("Contract (Word file only; open it and Save As PDF):", base + ".docx")

def to_pdf(docx_path, pdf_path):
    """Try LibreOffice, then Microsoft Word (via docx2pdf). Returns True if a PDF was made."""
    import shutil
    soffice = shutil.which("soffice") or shutil.which("libreoffice") or next(
        (p for p in (r"C:\Program Files\LibreOffice\program\soffice.exe",
                     r"C:\Program Files (x86)\LibreOffice\program\soffice.exe") if os.path.exists(p)), None)
    if soffice:
        try:
            subprocess.run([soffice, "--headless", "--convert-to", "pdf", "--outdir", OUT, docx_path],
                           check=True, capture_output=True, timeout=120)
            if os.path.exists(pdf_path): return True
        except Exception:
            pass
    try:
        from docx2pdf import convert      # needs Microsoft Word installed
        convert(docx_path, pdf_path)
        return os.path.exists(pdf_path)
    except Exception:
        return False

def main():
    p = argparse.ArgumentParser(description=__doc__); sub = p.add_subparsers(required=True)
    x = sub.add_parser("add-student"); x.set_defaults(f=cmd_add_student)
    x.add_argument("name"); x.add_argument("student_no"); x.add_argument("--grade")
    x.add_argument("--id-number", help="SA ID or passport number")
    x.add_argument("--guardian"); x.add_argument("--guardian-email")
    x = sub.add_parser("import-students", help="bulk add from CSV"); x.set_defaults(f=cmd_import); x.add_argument("file")
    x = sub.add_parser("import-laptops", help="bulk add laptops from CSV"); x.set_defaults(f=cmd_import_laptops); x.add_argument("file")
    x = sub.add_parser("add-laptop"); x.set_defaults(f=cmd_add_laptop)
    x.add_argument("tag"); x.add_argument("--model"); x.add_argument("--serial"); x.add_argument("--value", type=float, default=0)
    x = sub.add_parser("assign", help="assign laptop and make contract"); x.set_defaults(f=cmd_assign)
    x.add_argument("student_no"); x.add_argument("tag"); x.add_argument("--date")
    x = sub.add_parser("return"); x.set_defaults(f=cmd_return); x.add_argument("tag"); x.add_argument("--date")
    x = sub.add_parser("signed", help="mark contract as signed"); x.set_defaults(f=cmd_signed)
    x.add_argument("id", type=int); x.add_argument("--date")
    x = sub.add_parser("list"); x.set_defaults(f=cmd_list); x.add_argument("--all", action="store_true")
    x = sub.add_parser("available"); x.set_defaults(f=cmd_available)
    x = sub.add_parser("contract", help="regenerate a contract"); x.set_defaults(f=cmd_contract); x.add_argument("id", type=int)
    a = p.parse_args(); a.f(a)

if __name__ == "__main__": main()
